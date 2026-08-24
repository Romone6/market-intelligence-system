"""Export the blind Stage 7 calibration queue as a reviewable Word document."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


QUEUE_HASH = "5cc1ff2c0e0b182cc5d568d2469300d69c27dfba1010726153de69d6f6f0736f"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "20252B"
MUTED = "667085"
PALE_BLUE = "E8EEF5"
RULE = "CDD5DF"


def _xml_safe(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(
        r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]",
        " ",
        text,
    )


def _set_cell_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def _set_bottom_border(paragraph, color: str = RULE, size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = document.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_after = Pt(8)

    subtitle = document.styles.add_style("Corpus Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(BLUE)
    subtitle.paragraph_format.space_after = Pt(24)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    meta = document.styles.add_style("Corpus Metadata", WD_STYLE_TYPE.PARAGRAPH)
    meta.font.name = "Calibri"
    meta.font.size = Pt(9)
    meta.font.color.rgb = RGBColor.from_string(MUTED)
    meta.paragraph_format.space_after = Pt(3)

    source = document.styles.add_style("Source Text", WD_STYLE_TYPE.PARAGRAPH)
    source.font.name = "Calibri"
    source.font.size = Pt(9)
    source.font.color.rgb = RGBColor.from_string(INK)
    source.paragraph_format.space_after = Pt(6)
    source.paragraph_format.line_spacing = 1.08
    source.paragraph_format.widow_control = True


def _add_labelled_value(document: Document, label: str, value: object) -> None:
    paragraph = document.add_paragraph(style="Corpus Metadata")
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    paragraph.add_run(_xml_safe(value) if value not in (None, "") else "None supplied")


def _add_source_text(document: Document, text: object) -> None:
    safe = _xml_safe(text).strip()
    if not safe:
        document.add_paragraph("None supplied", style="Source Text")
        return
    blocks = re.split(r"\n\s*\n", safe)
    for block in blocks:
        paragraph = document.add_paragraph(style="Source Text")
        paragraph.add_run(block.strip())


def _load_tasks(database_path: Path) -> list[dict[str, object]]:
    with sqlite3.connect(database_path) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT event_id, company, entity_id, related_event_group,
                   source_published_at, tradable_at, filing_type, payload_json
            FROM annotation_tasks
            WHERE annotation_round = 'calibration'
            ORDER BY event_id
            """
        ).fetchall()
    tasks: list[dict[str, object]] = []
    for row in rows:
        payload = json.loads(row["payload_json"])
        tasks.append(
            {
                "event_id": row["event_id"],
                "company": row["company"],
                "entity_id": row["entity_id"],
                "related_event_group": row["related_event_group"],
                "source_published_at": row["source_published_at"],
                "tradable_at": row["tradable_at"],
                "filing_type": row["filing_type"],
                "prior_company_disclosure": payload.get("prior_company_disclosure"),
                "relevant_section": payload.get("relevant_section", ""),
            }
        )
    return tasks


def build_document(database_path: Path, output_path: Path) -> Path:
    tasks = _load_tasks(database_path)
    if len(tasks) != 100:
        raise ValueError(f"expected exactly 100 calibration tasks; found {len(tasks)}")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_styles(document)

    core = document.core_properties
    core.title = "Financial Event Calibration Corpus — 100 Blind SEC Review Tasks"
    core.subject = "Stage 7 calibration source packet"
    core.author = "QUANI HOOTT"
    core.keywords = "financial events, SEC filings, annotation, calibration"

    accent = document.add_paragraph()
    accent.paragraph_format.space_after = Pt(28)
    accent_run = accent.add_run("QUANI HOOTT  /  STAGE 7")
    accent_run.bold = True
    accent_run.font.name = "Calibri"
    accent_run.font.size = Pt(9)
    accent_run.font.color.rgb = RGBColor.from_string(BLUE)
    _set_bottom_border(accent, BLUE, "18")

    document.add_paragraph("Financial Event\nCalibration Corpus", style="Title")
    document.add_paragraph("100 blind SEC review tasks", style="Corpus Subtitle")

    cover_note = document.add_paragraph()
    cover_note.paragraph_format.space_after = Pt(22)
    cover_note.paragraph_format.line_spacing = 1.2
    run = cover_note.add_run(
        "A source-only review packet generated from the project’s fixed calibration queue. "
        "Candidate labels and machine-selected evidence highlights are intentionally excluded."
    )
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(INK)
    _set_cell_shading(cover_note, PALE_BLUE)

    _add_labelled_value(document, "Tasks", len(tasks))
    _add_labelled_value(document, "Distinct entities", len({task["entity_id"] for task in tasks}))
    _add_labelled_value(document, "Queue hash", QUEUE_HASH)
    _add_labelled_value(document, "Generated", datetime.now().astimezone().strftime("%d %B %Y %H:%M %Z"))
    _add_labelled_value(document, "Source database", database_path.resolve())

    document.add_page_break()

    for index, task in enumerate(tasks, start=1):
        if index > 1:
            document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        heading.add_run(f"Task {index:03d} of 100 — {_xml_safe(task['company'])}")
        _set_bottom_border(heading)

        _add_labelled_value(document, "Task ID", task["event_id"])
        _add_labelled_value(document, "Entity", task["entity_id"])
        _add_labelled_value(document, "Related filing / event group", task["related_event_group"])
        _add_labelled_value(document, "Filing type", task["filing_type"])
        _add_labelled_value(document, "Published", task["source_published_at"])
        _add_labelled_value(document, "First tradable", task["tradable_at"])

        document.add_paragraph("Review text", style="Heading 2")
        _add_source_text(document, task["relevant_section"])

        document.add_paragraph("Prior company disclosure", style="Heading 2")
        prior = task["prior_company_disclosure"]
        if prior:
            prior_note = document.add_paragraph(
                "Earlier disclosure supplied by the queue for novelty comparison.",
                style="Corpus Metadata",
            )
            prior_note.paragraph_format.space_after = Pt(6)
            _add_source_text(document, prior)
        else:
            document.add_paragraph("None supplied", style="Source Text")

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "FINANCIAL EVENT CALIBRATION CORPUS  ·  BLIND REVIEW PACKET"
        header.style = document.styles["Corpus Metadata"]
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_bottom_border(header, RULE, "4")
        _add_page_number(section.footer.paragraphs[0])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    digest = hashlib.sha256(output_path.read_bytes()).hexdigest()
    print(f"created={output_path.resolve()}")
    print(f"tasks={len(tasks)}")
    print(f"sha256={digest}")
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--database",
        type=Path,
        default=Path("data/labels/annotations.sqlite"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("reports/financial_event_calibration_corpus_100.docx"),
    )
    args = parser.parse_args()
    build_document(args.database, args.output)


if __name__ == "__main__":
    main()
